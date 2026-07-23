"""Workbook inspection tests: structural extraction from a real openpyxl-built
workbook, and the pure VBA procedure/trigger scanner (no oletools needed).

Real .xlsx files are built in-test with openpyxl. Extracting VBA from a genuine
vbaProject.bin is covered by the parser's own tests; here the new logic — the
structural inspection and the auto-run trigger detection — is what is exercised.
"""

from __future__ import annotations

import os
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table

from app.spreadsheet import inspect_file
from app.spreadsheet.inspect import scan_vba_procedures
from app.spreadsheet.models import VbaModule


def _sample_workbook() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Region", "Q1", "Q2", "Total"])
    ws.append(["East", 10, 20, "=B2+C2"])
    ws.append(["West", 5, 7, "=B3+C3"])
    ws.add_table(Table(displayName="SalesTbl", ref="A1:D3"))
    wb.defined_names.add(DefinedName("TaxRate", attr_text="Data!$B$1"))
    hidden = wb.create_sheet("Scratch")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "internal"
    path = Path(tempfile.mktemp(suffix=".xlsx"))
    wb.save(path)
    return path


def test_inspect_structure_reports_sheets_names_tables_formulas():
    path = _sample_workbook()
    try:
        insp = inspect_file(path, "budget.xlsx", size_bytes=os.path.getsize(path))
    finally:
        os.unlink(path)

    assert insp.file_format == "xlsx"
    assert insp.macro_enabled is False
    assert insp.has_macros is False
    assert insp.sheet_count == 2

    by_name = {s.name: s for s in insp.sheets}
    assert set(by_name) == {"Data", "Scratch"}
    assert by_name["Data"].formula_count == 2
    assert by_name["Data"].table_count == 1
    assert by_name["Scratch"].visibility == "hidden"

    names = {n.name: n for n in insp.named_ranges}
    assert "TaxRate" in names
    assert "Data" in names["TaxRate"].refers_to

    tables = {t.name: t for t in insp.tables}
    assert "SalesTbl" in tables
    assert tables["SalesTbl"].sheet == "Data"
    assert tables["SalesTbl"].columns == ["Region", "Q1", "Q2", "Total"]

    # No pivots/connections/charts in this workbook — must be empty, not crash.
    assert insp.pivot_tables == []
    assert insp.connections == []


def test_inspect_json_uses_camelcase_keys_the_extension_reads():
    path = _sample_workbook()
    try:
        payload = inspect_file(path, "budget.xlsx").model_dump(by_alias=True)
    finally:
        os.unlink(path)
    for key in ("fileFormat", "macroEnabled", "hasMacros", "sheetCount",
                "vbaModules", "namedRanges", "pivotTables", "aiContext",
                "aiContextTruncated"):
        assert key in payload, key

    # The prompt-ready context is richer than the visual sheet summary: it
    # includes extracted formulas and sample cell values for the browser agent.
    assert "SalesTbl" in payload["aiContext"]
    assert "=B2+C2" in payload["aiContext"]
    assert "East" in payload["aiContext"]
    assert payload["aiContextTruncated"] is False


def test_scan_vba_detects_procedures_and_auto_run_triggers():
    code = (
        "Private Sub Workbook_Open()\n"
        "    Run\n"
        "End Sub\n"
        "Sub Auto_Open()\n"
        "End Sub\n"
        "Public Function CalcTax(amt As Double) As Double\n"
        "    CalcTax = amt * 0.2\n"
        "End Function\n"
        "Private Sub Worksheet_Change(ByVal Target As Range)\n"
        "End Sub\n"
    )
    procs, triggers = scan_vba_procedures([VbaModule(name="ThisWorkbook", line_count=10, code=code)])

    by_name = {p.name: p for p in procs}
    assert set(by_name) == {"Workbook_Open", "Auto_Open", "CalcTax", "Worksheet_Change"}
    assert by_name["CalcTax"].kind == "Function"
    assert by_name["CalcTax"].auto_run is False
    assert by_name["Workbook_Open"].auto_run is True
    assert by_name["Auto_Open"].auto_run is True

    trig = {t.name: t for t in triggers}
    assert trig["Auto_Open"].scope == "auto"
    assert trig["Workbook_Open"].scope == "event"
    assert trig["Worksheet_Change"].scope == "event"
    assert "CalcTax" not in trig


def test_scan_vba_handles_empty_and_property_procs():
    code = (
        "Property Get Rate() As Double\n"
        "End Property\n"
        "Property Let Rate(v As Double)\n"
        "End Property\n"
    )
    procs, triggers = scan_vba_procedures([VbaModule(name="Cls", code=code)])
    kinds = sorted(p.kind for p in procs)
    assert kinds == ["Property Get", "Property Let"]
    assert triggers == []
    # No modules → no procedures, no crash.
    assert scan_vba_procedures([]) == ([], [])
